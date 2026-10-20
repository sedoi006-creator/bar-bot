# -*- coding: utf-8 -*-
"""Парсер технологических карт из Excel -> recipes.json для Telegram-бота барменов."""
import json
import re
import unicodedata
from pathlib import Path

import openpyxl

XLSX = Path(r"C:\Users\serj\Desktop\апхыпхыл\бар\концепт меню\Технологички БАР  (АвтоматическиВосстановлено).xlsx")
OUT = Path(__file__).parent / "recipes.json"

# Лист -> (категория в боте, режим разбора)
SHEETS = {
    "коктейли апарпыл": ("Коктейли", "table"),
    "Коктейли Доп": ("Коктейли", "table"),
    "Коктейли буд": ("Коктейли", "table"),
    "Лимонады": ("Лимонады", "table"),
    "Смузи": ("Смузи", "table"),
    " милшейки буд ": ("Милкшейки", "table"),
    "Milk Shake": ("Милкшейки", "table"),
    "Морс & Компот ": ("Морсы и компоты", "table"),
    "Чай": ("Чай", "table"),
    "Кофе": ("Кофе", "table"),
    "Наливки": ("Наливки", "table"),
    "Полуфабрикаты": ("Полуфабрикаты", "table"),
    "Фрукты пф": ("Полуфабрикаты", "table"),
    "чай заварной": ("Чай", "brewed_tea"),
    "пиво разливное": ("Пиво", "table"),
    "моктейли": ("Моктейли", "numbered"),
}

# Сопоставление заголовков колонок -> внутренние поля
COLMAP = {
    "название": "name",
    "компоненты": "component",
    "ингредиенты": "component",
    "мл/гр": "amount",
    "мл\\гр": "amount",
    "гр/мл": "amount",
    "метод": "method",
    "метод приготовления": "method",
    "лёд": "ice",
    "лед": "ice",
    "бокал": "glass",
    "украшение": "garnish",
    "укращение": "garnish",
    "подача": "serving",
    "выход": "output",
    "выход (мл)": "output",
    "выход мл": "output",
    "технология": "tech",
    "приготовление": "tech",
    "комментарии": "comment",
}


# Полные ручные карточки (перезаписывают данные из Excel). Ключ — название в нижнем регистре.
MANUAL_CARDS = {
    "кордиал реган": {
        "name": "Кордиал Реган (п/ф)",
        "category": "Полуфабрикаты",
        "components": [
            {"name": "Листья красного базилика (без стеблей)", "amount": "75 г"},
            {"name": "Суперджус лимон (п/ф)", "amount": "170 мл"},
            {"name": "Лимонная кислота", "amount": "27 г"},
            {"name": "Аскорбиновая кислота", "amount": "3 г"},
            {"name": "Вода холодная (питьевая)", "amount": "769 мл"},
            {"name": "Сахар", "amount": "405 г"},
            {"name": "Агар-агар", "amount": "3 г"},
        ],
        "method": "Блендер → кипение 1/4 с агаром → застывание → отжим на марле",
        "ice": "", "glass": "", "garnish": "", "serving": "",
        "output": "1000 мл",
        "comment": "Хранить при +2…+6 °C не более 5–7 суток. Перед использованием — визуальный контроль: при осадке, помутнении или постороннем запахе — утилизировать.",
        "tech": (
            "1) Подготовка сырья. Листья красного базилика перебрать, удалить стебли.\n"
            "2) Смешивание основы. В блендер заложить листья базилика, суперджус лимон, лимонную "
            "кислоту, аскорбиновую кислоту, холодную воду и сахар. Пробить до однородной массы.\n"
            "3) Желирующая часть. Отобрать 1/4 полученной массы, добавить агар-агар. Довести до "
            "кипения при постоянном помешивании — до появления крупных пузырей.\n"
            "4) Соединение. Горячую часть с агаром смешать с оставшейся холодной массой. "
            "Тщательно перемешать.\n"
            "5) Застывание. Перелить массу в шейкер (или ёмкость) и убрать до полного застывания.\n"
            "6) Отделение жидкости. После застывания откинуть массу на марлю (в несколько слоёв) "
            "и дать полностью откапаться. Отделившаяся жидкость — это и есть кордиал.\n"
            "7) Розлив. Перелить кордиал в стерильную стеклянную тару, промаркировать "
            "(название, дата, партия)."
        ),
    },
}

# Названия, утерянные при восстановлении Excel-файла, восстановлены вручную.
# Ключ — имя, составленное парсером из первых трёх ингредиентов.
ANON_RENAMES = {
    "персиковый ликер + кордиал жасмин + фруктовое игристое вино": "Peach Fiz",
    "цитрусовая водка + розовая вода + пюре клубники": "Rossini",
    "лимончелло + суперджус апельсин + аромка бергамот": "Limoncello Spritz",
    "текила + кампари + арбузный кордиал": "Paloma Арбуз",
    "черный ром + суперджус апельсин + кордиал фалернум": "Медовуха Спритц",
    "джин + пюре маракуйя + суперджус лимон": "Bogini",
    "джим бим + амаретто + суперджус лимон": "Almond Sour",
    "абсолют груша + мед + мята": "Julep Mule",
    "виски + угольный краситель + кордиал ройбуш юдзу": "Sparcle Peniciline",
    "джин + ликер бузина + пюре личи": "Elderflower Lady",
    "бренди + мараскино ликер + суперджус лимон": "Brandy Special",
    "виски + вермут + черрри бренди": "Янтарный",
    "Водка + джин 20 + вермут сухой": "Веспер",
    "водка 20 + виски айла + какао биттер": "Белый Русский",
    "ром + мятный ликер + сливочная фейхоа": "Grasshopper",
}


def norm(s):
    if s is None:
        return ""
    s = str(s).replace("\xa0", " ").replace("\n", " ")
    s = unicodedata.normalize("NFKC", s)
    return re.sub(r"\s+", " ", s).strip()


def header_fields(row):
    """Если строка — заголовок («Название» ...), вернуть {колонка: поле}."""
    fields = {}
    for i, cell in enumerate(row):
        key = norm(cell).lower()
        if key in COLMAP:
            fields[i] = COLMAP[key]
    if fields.get(0) == "name":
        return fields
    return None


def parse_table(ws, skip_rows=None):
    """Универсальный разбор листа-таблицы: заголовок -> блок рецепта.
    skip_rows — диапазон (начало, конец) строк, которые разбирает другой парсер."""
    drinks = []
    fields = None
    current = None

    for row_num, row in enumerate(ws.iter_rows(values_only=True), 1):
        if skip_rows and skip_rows[0] <= row_num <= skip_rows[1]:
            continue
        cells = [norm(c) for c in row]
        if not any(cells):
            continue
        hdr = header_fields(cells)
        if hdr:
            fields = hdr
            current = None
            continue
        if fields is None:
            continue

        first = cells[0]
        if first:
            # Новая карточка: первая ячейка — название (или служебная пометка)
            if first.lower().startswith("выход") and current:
                current["output"] = first
                continue
            if first.lower() == "название":
                continue
            current = {"name": first, "components": [], "_f": {k: [] for k in
                      ("method", "ice", "glass", "garnish", "serving", "output", "tech", "comment")}}
            drinks.append(current)
        if current is None:
            # Блок без названия (название потеряно при восстановлении файла) —
            # соберём компоненты, имя присвоим ниже по первым ингредиентам
            comp_check = [cells[i] for i in fields if i < len(cells) and cells[i]
                          and fields[i] in ("component", "amount")]
            if not comp_check:
                continue
            current = {"name": "", "components": [], "_f": {k: [] for k in
                      ("method", "ice", "glass", "garnish", "serving", "output", "tech", "comment")},
                       "_anon": True}
            drinks.append(current)

        comp_name, comp_amount = "", ""
        for i, val in enumerate(cells[1:], 1):
            if not val:
                continue
            field = fields.get(i, "component")
            if field == "component":
                comp_name = val if not comp_name else f"{comp_name}; {val}"
            elif field == "amount":
                comp_amount = val
            else:
                bucket = current["_f"].setdefault(field, [])
                if val not in bucket:
                    bucket.append(val)
        if comp_name:
            current["components"].append({"name": comp_name, "amount": comp_amount})
    return drinks


def parse_multicolumn(ws, start_row, end_row):
    """Зона «шприц-коктейлей» (строки 38–60 листа апарпыл): 3 колонки на рецепт.
    Внутри колонки: «N. Название», компоненты, «Украшение:», «Бокал:», «Метод:»."""
    drinks = []
    rows = list(ws.iter_rows(min_row=start_row, max_row=end_row, values_only=True))
    for col0 in (0, 3, 6, 9):  # стартовые колонки блоков (A, D, G, J)
        current = None
        for row in rows:
            cells = [norm(c) for c in row[col0:col0 + 3]]
            first = cells[0] if cells else ""
            if not first:
                continue
            m = re.match(r"^\d+\s*[.)]\s*(.+)", first)
            if m:
                current = {"name": m.group(1).strip(), "components": [], "_f": {}}
                drinks.append(current)
                # на строке с названием иногда уже лежит компонент во 2-й колонке
                if len(cells) > 1 and cells[1] and not cells[1].startswith(
                        ("Украшение", "Бокал", "Метод")):
                    current["components"].append({"name": cells[1], "amount": cells[2]})
                continue
            if current is None:
                continue
            low = first.lower()
            if low.startswith(("украшение", "укращение")):
                current["_f"].setdefault("garnish", []).append(
                    first.split(":", 1)[-1].strip() or first)
            elif low.startswith("бокал"):
                current["_f"].setdefault("glass", []).append(
                    first.split(":", 1)[-1].strip() or first)
            elif low.startswith("метод"):
                current["_f"].setdefault("method", []).append(
                    first.split(":", 1)[-1].strip() or first)
            else:
                current["components"].append({"name": first, "amount": cells[1] if len(cells) > 1 else ""})
    return drinks


def parse_numbered(ws):
    """Лист «моктейли»: «1. Название» -> компоненты в 2 колонках."""
    drinks, current = [], None
    for row in ws.iter_rows(values_only=True):
        cells = [norm(c) for c in row]
        first = cells[0] if cells else ""
        second = cells[1] if len(cells) > 1 else ""
        if not first:
            continue
        if re.match(r"^\d+\s*[.)]", first):
            current = {"name": re.sub(r"^\d+\s*[.)]\s*", "", first),
                       "components": [], "_f": {}}
            drinks.append(current)
            continue
        if current is None:
            current = {"name": first, "components": [], "_f": {}}
            drinks.append(current)
            continue
        if first.lower().startswith(("украш", "укращ")):
            current["_f"]["garnish"] = [first]
        else:
            current["components"].append({"name": first, "amount": second})
    return drinks


def parse_brewed_tea(ws):
    """Лист «чай заварной»: простая таблица название / гр / выход."""
    drinks = []
    for row in ws.iter_rows(values_only=True):
        cells = [norm(c) for c in row]
        if not cells or not cells[0] or cells[0].lower() == "название":
            continue
        if "черные чаи" in cells[0].lower():
            continue
        drink = {"name": cells[0], "components": [], "_f": {}}
        if len(cells) > 1 and cells[1]:
            drink["components"].append({"name": "Чай сухой", "amount": cells[1] + " гр"})
        if len(cells) > 2 and cells[2]:
            drink["_f"]["output"] = cells[2] + " мл"
        drinks.append(drink)
    return drinks


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    all_drinks = []
    for sheet_name, (category, mode) in SHEETS.items():
        if sheet_name not in wb.sheetnames:
            print(f"!! Лист не найден: {sheet_name!r}")
            continue
        ws = wb[sheet_name]
        if mode == "table":
            skip = (38, 60) if sheet_name == "коктейли апарпыл" else None
            raw = parse_table(ws, skip_rows=skip)
            if sheet_name == "коктейли апарпыл":
                raw += parse_multicolumn(ws, 38, 60)
        elif mode == "numbered":
            raw = parse_numbered(ws)
        else:
            raw = parse_brewed_tea(ws)

        seen = set()
        count = 0
        for d in raw:
            name = d["name"]
            if d.pop("_anon", False) and not name:
                # имя потеряно при восстановлении файла — составим из ингредиентов
                parts = [c["name"] for c in d["components"][:3]]
                name = " + ".join(parts) if parts else ""
                name = ANON_RENAMES.get(name, name)
            if not name or len(name) < 2:
                continue
            # отсекаем обрывки технологий, попавшие в колонку названий
            if re.match(r"^[\d.,;]+\s*(этап|шаг)\b", name, re.IGNORECASE):
                continue
            # убираем дубли внутри листа и шум восстановления файла
            key = name.lower().strip(" .")
            if key in seen:
                continue
            seen.add(key)
            f = d.pop("_f", {})
            drink = {
                "name": name,
                "category": category,
                "sheet": sheet_name.strip(),
                "components": d["components"],
                "method": " / ".join(f.get("method", [])),
                "ice": " / ".join(f.get("ice", [])),
                "glass": " / ".join(f.get("glass", [])),
                "garnish": " / ".join(f.get("garnish", [])),
                "serving": " / ".join(f.get("serving", [])),
                "output": " / ".join(f.get("output", [])),
                "tech": " ".join(f.get("tech", [])),
                "comment": " / ".join(f.get("comment", [])),
            }
            # пропускаем пустышки без единого компонента и технологии
            if not drink["components"] and not drink["tech"] and not drink["serving"]:
                continue
            all_drinks.append(drink)
            count += 1
        print(f"{sheet_name!r:25s} -> {category:15s} карточек: {count}")

    # финальная дедупликация по имени+категории (например, «манго маракуйя» дважды)
    final, keys = [], set()
    for d in all_drinks:
        key = (d["name"].lower(), d["category"])
        if key in keys:
            # объединяем с уже добавленным
            for prev in final:
                if (prev["name"].lower(), prev["category"]) == key:
                    prev["components"] = prev["components"] or d["components"]
                    for fld in ("method", "ice", "glass", "garnish", "serving", "output", "tech"):
                        if not prev[fld] and d[fld]:
                            prev[fld] = d[fld]
                    break
            continue
        keys.add(key)
        final.append(d)

    # применяем полные ручные карточки (перезапись или добавление)
    for mkey, card in MANUAL_CARDS.items():
        replaced = False
        for i, d in enumerate(final):
            if d["name"].lower() == mkey:
                card = {**card, "sheet": "вручную"}
                final[i] = card
                replaced = True
                break
        if not replaced:
            final.append({**card, "sheet": "вручную"})

    OUT.write_text(json.dumps(final, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"\nИТОГО рецептов: {len(final)} -> {OUT}")
    cats = {}
    for d in final:
        cats[d["category"]] = cats.get(d["category"], 0) + 1
    for c, n in sorted(cats.items()):
        print(f"  {c}: {n}")


if __name__ == "__main__":
    main()
